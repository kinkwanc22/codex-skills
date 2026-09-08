from pathlib import Path
import re,json,hashlib
from docx import Document
from docx.shared import Pt,Cm
from docx.oxml.ns import qn

import argparse
ap=argparse.ArgumentParser();ap.add_argument('--run-dir',type=Path,required=True);ap.add_argument('--output',type=Path,required=True);ap.add_argument('--body',type=Path);args=ap.parse_args()
root=args.run_dir.resolve()
raw=(args.body or root/'expanded_raw.txt').read_text().strip()
review=json.loads((root/'final_dedup_review.json').read_text())
assert review.get('pass') is True and review.get('compared_ids') and review.get('findings'),'final semantic dedup review required'
assert review['body_sha256']==hashlib.sha256(raw.encode()).hexdigest(),'body changed after review'
(root/'delivery_body.txt').write_text(raw)
out=args.output.resolve()
out.parent.mkdir(parents=True,exist_ok=True)
d=Document()
s=d.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=s.bottom_margin=Cm(2)
s.left_margin=s.right_margin=Cm(2.2)
normal=d.styles['Normal']
normal.font.name='Arial Unicode MS';normal.font.size=Pt(11)
normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
normal.paragraph_format.line_spacing=1.35
normal.paragraph_format.space_after=Pt(8)
for block in raw.split('\n'):
    if block.strip():d.add_paragraph(block)
d.save(out)
back='\n'.join(p.text for p in Document(out).paragraphs)
expected='\n'.join(x for x in raw.split('\n') if x.strip())
assert back==expected
report={'output':str(out),'body_cjk':len(re.findall(r'[\u3400-\u9fff]',raw)),'docx_text_matches_delivery':True,'length_gate':False,'source_sha256':hashlib.sha256((root/'approved_source.txt').read_bytes()).hexdigest(),'input_mode':'new source plus case instruction into fresh old25 clone','state':'正文待确认'}
(root/'delivery_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False))
